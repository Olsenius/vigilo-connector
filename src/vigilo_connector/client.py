"""HTTP-klient mot Vigilos foreldre-API.

Endepunktene under `# Kjente endepunkter` er verifisert via
https://github.com/brujoand/vigilo2smtp. Nye endepunkter (ukesplan, nyheter,
fravær) legges til her etter hvert som de kartlegges — bruk `get()` /
MCP-verktøyet `api_get` til å utforske.
"""

import io
import math
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed, TimeoutError as FuturesTimeout
from contextlib import contextmanager
import re
from datetime import date, datetime, time as datetime_time, timedelta
from zoneinfo import ZoneInfo
from urllib.parse import quote

import httpx

from .auth import AuthError, AuthLockTimeout, TokenStore, auth_budget, user_id_from_jwt
from .http_logging import LoggedClient, logged_get
from .config import API_BASE, APP_VERSION, WEB_API_BASE
from .transport import JSON_CONTENT_TYPE, android_client, android_json, is_filestorage


class WriteOutcomeUnknown(RuntimeError):
    """Skrivekallet kan ha blitt behandlet; ikke send det på nytt automatisk."""


class SearchIncomplete(RuntimeError):
    """Forventet ufullstendig søk; trygt å vise uten serverinnhold."""


def _oslo_today() -> date:
    return datetime.now(ZoneInfo("Europe/Oslo")).date()


def _search_window(since: str | None) -> tuple[str, datetime]:
    if since is None:
        today = _oslo_today()
        start = date(today.year if today.month >= 8 else today.year - 1, 8, 1)
    else:
        if not isinstance(since, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", since):
            raise ValueError("since må vere ein ISO-dato: YYYY-MM-DD.")
        try:
            start = date.fromisoformat(since)
        except ValueError:
            raise ValueError("since må vere ein gyldig ISO-dato: YYYY-MM-DD.") from None
    return start.isoformat(), datetime.combine(start, datetime_time.min, ZoneInfo("Europe/Oslo"))


def _message_time(value) -> datetime:
    # fromisoformat normaliserer til dømes +00:60; kontroller offset før parsing.
    if not isinstance(value, str):
        raise ValueError("Manglande timeCreated.")
    match = re.fullmatch(
        r"\d{4}-\d{2}-\d{2}[Tt]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?P<offset>[Zz]|[+-]\d{2}:\d{2})",
        value,
    )
    if match is None:
        raise ValueError("timeCreated må ha dato, klokkeslett og eksplisitt tidssone.")
    offset = match.group("offset")
    if offset.upper() != "Z" and (int(offset[1:3]) > 23 or int(offset[4:6]) > 59):
        raise ValueError("Ugyldig tidssone i timeCreated.")
    return datetime.fromisoformat(value.replace("t", "T").replace("z", "Z"))


def _absence_dates(from_date: str, to_date: str) -> tuple[str, str]:
    """ISO-dager inkludert begge ender til appens eksklusive datoperiode."""
    if not all(isinstance(d, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", d)
               for d in (from_date, to_date)):
        raise ValueError("Datoene må ha format YYYY-MM-DD.")
    start, end = date.fromisoformat(from_date), date.fromisoformat(to_date)
    if end < start:
        raise ValueError("to_date kan ikke være før from_date.")
    if end == date.max:
        raise ValueError("to_date må tillate sluttdato dagen etter.")
    return start.strftime("%d.%m.%Y"), (end + timedelta(days=1)).strftime("%d.%m.%Y")


def iso_week(d: date | None = None) -> str:
    """ISO-uke på formatet Vigilo bruker, f.eks. "2026-38"."""
    y, w, _ = (d or date.today()).isocalendar()
    return f"{y}-{w:02d}"


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(p for p in parts if p).strip()


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _attachment_to_text(data: bytes, mime: str, name: str | None) -> str | None:
    """Trekk ut tekst fra et vedlegg, eller None for formater vi ikke leser."""
    n = (name or "").lower()
    if mime == "application/pdf" or n.endswith(".pdf"):
        return _pdf_to_text(data)
    if "wordprocessingml" in mime or n.endswith(".docx"):
        return _docx_to_text(data)
    if mime.startswith("text/"):
        return data.decode("utf-8", "replace")
    return None


class VigiloClient:
    def __init__(self, store: TokenStore | None = None):
        self._read_context = threading.local()
        self._store = store or TokenStore()
        self._http = android_client(API_BASE, app_api=True)
        self._web = LoggedClient(origin="web", base_url=WEB_API_BASE, timeout=30)

    def _headers(self, force_refresh: bool = False) -> dict:
        return {
            "Authorization": f"Bearer {self._access_token(force_refresh)}",
            "appVersion": APP_VERSION,
        }

    def _access_token(self, force_refresh=False):
        budget = getattr(self._read_context, "budget", None)
        if budget is None:
            return self._store.access_token(force_refresh)
        deadline, stopped = budget
        self._read_options()
        previous_phase = getattr(self._read_context, "phase", "scan")
        self._read_context.phase = "auth"
        try:
            with auth_budget(min(deadline, time.monotonic() + 15.0)):
                token = self._store.access_token(force_refresh)
        except AuthLockTimeout:
            self._read_context.phase = "auth_lock"
            raise
        self._read_options()
        self._read_context.phase = previous_phase
        return token

    def _search_diagnostics(self, started, default_phase):
        phase = getattr(self._read_context, "phase", default_phase)
        return f"phase={phase}, elapsed_s={time.monotonic() - started:.3f}; "

    def _web_headers(self, force_refresh: bool = False) -> dict:
        # Web-API-et fester bare Authorization-header; ingen appVersion/userId.
        return {"Authorization": f"Bearer {self._access_token(force_refresh)}"}

    @property
    def user_id(self) -> str:
        return user_id_from_jwt(self._store.access_token())

    @contextmanager
    def _search_reads(self, deadline, stopped):
        previous = getattr(self._read_context, "budget", None)
        self._read_context.budget = (deadline, stopped)
        self._read_context.phase = "scan"
        try:
            yield
        finally:
            self._read_context.budget = previous
            self._read_context.first_list = False

    def _read_options(self):
        budget = getattr(self._read_context, "budget", None)
        if budget is None:
            return {}
        deadline, stopped = budget
        remaining = deadline - time.monotonic()
        if stopped.is_set() or remaining <= 0:
            raise SearchIncomplete("Tidsbudsjettet er brukt opp; søket er ikkje komplett.")
        if getattr(self._read_context, "first_list", False):
            connect = min(5.0, remaining / 4)
            return {"timeout": httpx.Timeout(
                connect=connect, read=min(15.0, remaining - connect),
                write=min(1.0, remaining / 20), pool=min(1.0, remaining / 20))}
        return {"timeout": httpx.Timeout(
            connect=min(5.0, remaining), read=min(15.0, remaining),
            write=min(5.0, remaining), pool=min(5.0, remaining))}

    def get(self, path: str, params: dict | None = None) -> httpx.Response:
        """Rått GET-kall med auth og retry ved 401. Brukes også til kartlegging."""
        if is_filestorage(path):
            return self._http.get(path, params=params, **self._read_options())
        headers = self._headers()
        r = self._http.get(path, params=params, headers=headers, **self._read_options())
        if r.status_code == 401:
            headers = self._headers(force_refresh=True)
            r = self._http.get(path, params=params, headers=headers, **self._read_options())
        return r

    def get_json(self, path: str, params: dict | None = None):
        r = self.get(path, params=params)
        if r.status_code == 401:
            raise AuthError("Fortsatt 401 etter token-fornyelse — kjør `vigilo-login` på nytt.")
        r.raise_for_status()
        return r.json()

    def web_get(self, path: str, params: dict | None = None) -> httpx.Response:
        """Rått GET mot web-foreldreportalen med auth og retry ved 401."""
        r = self._web.get(path, params=params, headers=self._web_headers())
        if r.status_code == 401:
            r = self._web.get(path, params=params, headers=self._web_headers(force_refresh=True))
        return r

    def web_get_json(self, path: str, params: dict | None = None):
        r = self.web_get(path, params=params)
        if r.status_code == 401:
            raise AuthError(
                "401 fra web-API-et. Enten er tokenet utløpt (kjør `vigilo-login`), "
                "eller så aksepterer ikke web-API-et app-tokenet — se README."
            )
        r.raise_for_status()
        return r.json()

    # --- Kjente endepunkter ---------------------------------------------

    def _write_json(
        self, method: str, path: str, payload: dict | None = None, *,
        headers: dict[str, str] | None = None, retry_401: bool = False,
        outcome: str = "Handlingen kan være utført. Kontroller gjeldende tilstand",
    ) -> httpx.Response:
        """Skriv uten redirects/retries; 401-fornyelse velges per operasjon.

        None betyr ingen body; {} betyr et tomt JSON-objekt. Ekstra headers
        gjelder bare dette endepunktet og kan ikke overstyre autentisering.
        """
        if method not in {"POST", "PUT", "DELETE"}:
            raise ValueError("Skriveverb må være POST, PUT eller DELETE.")
        def request(force_refresh: bool = False) -> httpx.Response:
            request_headers = {k: v for k, v in (headers or {}).items()
                               if k.lower() not in {"authorization", "appversion"}}
            request_headers.update(self._headers(force_refresh))
            if payload is not None:
                request_headers["Content-Type"] = JSON_CONTENT_TYPE
            return self._http.request(
                method, path, content=android_json(payload) if payload is not None else None,
                headers=request_headers, follow_redirects=False,
            )
        try:
            response = request()
            if response.status_code == 401 and retry_401:
                response = request(force_refresh=True)
        except httpx.TransportError:
            raise WriteOutcomeUnknown(
                f"Ukjent resultat: {outcome} før et nytt forsøk; ikke send automatisk på nytt."
            ) from None
        correlation_id = response.headers.get("X-Correlation-Id", "")
        correlation = (f" Korrelasjons-ID: {correlation_id}."
                       if re.fullmatch(r"[A-Za-z0-9_.:-]{1,128}", correlation_id) else "")
        if response.status_code == 401:
            raise AuthError("401 fra Vigilo — kjør `vigilo-login` på nytt." + correlation)
        if 300 <= response.status_code < 400 or response.status_code >= 500:
            raise WriteOutcomeUnknown(
                f"Vigilo svarte HTTP {response.status_code}. {outcome} før et nytt forsøk; "
                "ikke send automatisk på nytt." + correlation
            )
        if not 200 <= response.status_code < 300:
            raise httpx.HTTPStatusError(
                f"Vigilo avviste skrivekallet med HTTP {response.status_code}." + correlation,
                request=response.request, response=response,
            )
        return response

    def _post_absence(self, path: str, payload: dict) -> None:
        self._write_json(
            "POST", path, payload, retry_401=True,
            outcome="Fraværet kan være registrert. Kontroller fraværsoversikten",
        )

    def absence_codes(self, organizational_unit_id: str) -> list:
        """Aktive fraværskoder for den valgte barnehagen."""
        if not organizational_unit_id.strip():
            raise ValueError("organizational_unit_id må være utfylt.")
        data = self.get_json("/api/absencecodes", params={
            "organizationalUnitIds": organizational_unit_id,
        })
        return [code for item in data["items"]
                if item["organizationalUnitId"] == organizational_unit_id
                for code in item.get("codes") or [] if not code.get("isDeleted")]

    def register_childcare_absence(
        self, child_id: str, organizational_unit_id: str, from_date: str,
        to_date: str, absence_code_id: str, note: str | None = None,
    ) -> dict:
        """Barnehagefravær med enhetens aktive kode og valgfri merknad."""
        for field, value in (("child_id", child_id), ("organizational_unit_id", organizational_unit_id),
                             ("absence_code_id", absence_code_id)):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{field} må være utfylt.")
        if note is not None and not isinstance(note, str):
            raise ValueError("note må være tekst eller null.")
        start, end = _absence_dates(from_date, to_date)
        selected = next((code for code in self.absence_codes(organizational_unit_id)
                         if code["id"] == absence_code_id), None)
        if selected is None:
            raise ValueError("Velg en aktiv fraværskode fra absence_codes for denne barnehagen.")
        payload = {
            "childId": child_id, "organizationalUnitId": organizational_unit_id,
            "fromDate": start, "toDate": end, "absenceCodeId": absence_code_id,
        }
        if note is not None:
            payload["note"] = note
        self._post_absence("/api/absences", payload)
        return {
            "status": "registered", "childId": child_id,
            "organizationalUnitId": organizational_unit_id,
            "fromDate": from_date, "toDate": to_date,
            "absenceCode": {"id": selected["id"], "name": selected["name"]}, "note": note,
        }

    def message_contacts(self, child_id: str, organizational_unit_id: str) -> dict:
        """Kontaktgrupper med eksplisitte mottakerverdier for skrivekall."""
        data = self.get_json("/api/messages/contact-list", params={
            "childId": child_id, "organizationalUnitId": organizational_unit_id,
        })
        # Bruk kontaktens id, ikke employeeId. Behold grupper og øvrige felt.
        for key, recipient_type in (
            ("legalGuardians", "legalGuardian"),
            ("relatedEmployees", "employee"), ("otherEmployees", "employee"),
        ):
            for contact in data.get(key) or []:
                contact["recipient"] = {"type": recipient_type, "externalId": contact["id"]}
        for group in data.get("communicationGroups") or []:
            for contact in group.get("employees") or []:
                contact["recipient"] = {"type": "employee", "externalId": contact["id"]}
        return data

    def register_student_absence(
        self, child_id: str, organizational_unit_id: str, from_date: str,
        to_date: str, note: str, title: str, recipients: list[dict[str, str]],
    ) -> dict:
        """Opprett skolefravær uten vedlegg. Ingen automatisk retry ved transportfeil."""
        for field, value in (("child_id", child_id), ("organizational_unit_id", organizational_unit_id),
                             ("note", note), ("title", title)):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{field} må være utfylt.")
        start, end = _absence_dates(from_date, to_date)
        if not recipients:
            raise ValueError("Velg minst én mottaker fra message_contacts.")
        selected = []
        for recipient in recipients:
            if (set(recipient) != {"type", "externalId"}
                    or recipient["type"] not in {"employee", "legalGuardian"}
                    or not isinstance(recipient["externalId"], str)
                    or not recipient["externalId"].strip()):
                raise ValueError("Mottakere må ha type employee/legalGuardian og externalId fra kontaktlisten.")
            selected.append(dict(recipient))
        payload = {
            "childId": child_id, "organizationalUnitId": organizational_unit_id,
            "fromDate": start, "toDate": end,
            "note": note, "title": title, "attachments": [], "recipients": selected,
        }
        self._post_absence("/api/student-absences", payload)
        return {
            "status": "registered", "childId": child_id,
            "organizationalUnitId": organizational_unit_id,
            "fromDate": from_date, "toDate": to_date, "title": title,
            "recipients": selected,
        }

    def get_children(self) -> list:
        data = self.get_json("/api/children")
        return data if isinstance(data, list) else data.get("items", [])

    def _post_message(self, path: str, payload: dict) -> None:
        self._write_json(
            "POST", path, payload, retry_401=True,
            outcome="Meldingen kan være sendt. Kontroller meldingstrådene",
        )

    def send_message(
        self, child_id: str, organizational_unit_id: str, title: str, body: str,
        recipients: list[dict[str, str]],
    ) -> dict:
        """Ny melding uten vedlegg, til eksplisitte kontakter fra valgt barn/enhet."""
        for field, value in (("child_id", child_id), ("organizational_unit_id", organizational_unit_id),
                             ("title", title), ("body", body)):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{field} må være utfylt.")
        if not isinstance(recipients, list) or not recipients:
            raise ValueError("Velg mottakere fra message_contacts, inkludert minst én ansatt.")
        selected = []
        for recipient in recipients:
            if (not isinstance(recipient, dict) or set(recipient) != {"type", "externalId"}
                    or recipient["type"] not in ("employee", "legalGuardian")
                    or not isinstance(recipient["externalId"], str)
                    or not recipient["externalId"].strip()):
                raise ValueError("Mottakere må ha type employee/legalGuardian og externalId fra kontaktlisten.")
            if not any(r["externalId"] == recipient["externalId"] for r in selected):
                selected.append(dict(recipient))
        if not any(r["type"] == "employee" for r in selected):
            raise ValueError("Minst én lærer/administrasjon må være mottaker.")
        contacts = self.message_contacts(child_id, organizational_unit_id)
        allowed = [c["recipient"] for key in ("relatedEmployees", "otherEmployees", "legalGuardians")
                   for c in contacts.get(key) or []]
        allowed.extend(c["recipient"] for g in contacts.get("communicationGroups") or []
                       for c in g.get("employees") or [])
        if any(r not in allowed for r in selected):
            raise ValueError("Velg mottakere fra kontaktlisten for dette barnet og denne enheten.")
        # Appens use-case legger barnet først; ContactTypeMapper serialiserer
        # CHILD som "client". Mapperen dedupliserer deretter alle på kontakt-ID.
        wire_recipients = [{"type": "client", "externalId": child_id}]
        wire_recipients.extend(r for r in selected if r["externalId"] != child_id)
        self._post_message("/api/messages", {
            "organizationalUnitId": organizational_unit_id, "title": title, "body": body,
            "attachments": [], "recipients": wire_recipients,
        })
        return {"status": "sent", "childId": child_id, "title": title, "recipients": selected}

    def reply_message(self, thread_uid: str, body: str) -> dict:
        """Svar i en eksisterende ordinær meldingstråd uten vedlegg."""
        for field, value in (("thread_uid", thread_uid), ("body", body)):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{field} må være utfylt.")
        if thread_uid in (".", ".."):
            raise ValueError("Ugyldig thread_uid.")
        thread = self.get_thread(thread_uid)
        # Bare ordinært detaljendepunkt kan brukes for svar her. Fravær har
        # egne endepunkter og krever en separat, kartlagt skrivekontrakt.
        if thread.get("threadUid") != thread_uid or "isReplyDisabled" not in thread:
            raise ValueError("Kun en bekreftet ordinær meldingstråd kan besvares.")
        if thread["isReplyDisabled"] is not False:
            raise ValueError("Svar er deaktivert i denne meldingstråden.")
        if thread.get("isDeleted") is True:
            raise ValueError("En slettet meldingstråd kan ikke besvares.")
        self._post_message(f"/api/messages/{quote(thread_uid, safe='')}/reply", {
            "body": body, "attachments": [],
        })
        return {"status": "sent", "threadUid": thread_uid}

    @staticmethod
    def _paging(page_size: int, page_id: int | None) -> dict:
        if type(page_size) is not int or page_size < 1:
            raise ValueError("page_size må vere eit positivt heiltal.")
        if page_id is not None and (type(page_id) is not int or page_id < 0):
            raise ValueError("page_id må vere eit ikkje-negativt heiltal.")
        return {"pageSize": page_size, **({"pageId": page_id} if page_id is not None else {})}

    def get_message_threads_page(
        self, child_id: str, page_size: int = 50, include_after_school: bool = True,
        page_id: int | None = None,
    ) -> dict:
        """En side med rå metadata, inkludert pageId og siste melding per tråd."""
        return self.get_json("/api/message-threads", params={
            "childIds": child_id, **self._paging(page_size, page_id),
            "unreadMessageThreadsCountOnly": "false",
            "includeAfterSchoolMessages": str(include_after_school).lower(),
        })

    def get_message_threads(
        self, child_id: str, page_size: int = 50, include_after_school: bool = True,
        page_id: int | None = None,
    ) -> list:
        return self.get_message_threads_page(
            child_id, page_size, include_after_school, page_id,
        ).get("messageThreads", [])

    def get_thread(
        self, thread_uid: str, page_size: int = 50, page_id: int | None = None,
    ) -> dict:
        return self.get_json(
            f"/api/messages/threads/{quote(thread_uid, safe='')}",
            params=self._paging(page_size, page_id),
        )

    @staticmethod
    def _search_timeout(timeout_s):
        if (isinstance(timeout_s, bool) or not isinstance(timeout_s, (int, float))
                or not math.isfinite(timeout_s) or timeout_s <= 0 or timeout_s > 120):
            raise ValueError("timeout_s må vere eit tal større enn 0 og høgst 120.")

    def search_message_thread(
        self, thread_uid: str, query: str, page_size: int = 50, max_pages: int = 100,
        timeout_s: float = 45.0,
    ) -> dict:
        """Søk alle detaljsider med et tidsbudsjett; endrer ikke lesestatus."""
        self._search_timeout(timeout_s)
        inherited = getattr(self._read_context, "budget", None)
        deadline, stopped = inherited or (time.monotonic() + timeout_s, threading.Event())
        with self._search_reads(deadline, stopped):
            return self._search_message_thread(thread_uid, query, page_size, max_pages, timeout_s)

    def _search_message_thread(self, thread_uid, query, page_size, max_pages, timeout_s, since_cutoff=None):
        if not isinstance(thread_uid, str) or not thread_uid.strip() or thread_uid in (".", ".."):
            raise ValueError("thread_uid må vere ein gyldig ID.")
        if not isinstance(query, str) or not query.strip():
            raise ValueError("query må vere utfylt.")
        self._paging(page_size, None)
        if type(max_pages) is not int or max_pages < 1:
            raise ValueError("max_pages må vere eit positivt heiltal.")
        started = time.monotonic()
        pages_fetched = messages_scanned = 0
        def incomplete(reason):
            return SearchIncomplete(
                f"{reason} Søket er ikkje komplett. max_pages={max_pages}, "
                f"page_size={page_size}, timeout_s={timeout_s}; "
                + self._search_diagnostics(started, "scan") +
                f"pagesFetched={pages_fetched}, messagesScanned={messages_scanned}.")
        query_folded = query.casefold()
        page_id = None
        seen_pages, seen_messages = set(), set()
        matches = []
        for _ in range(max_pages):
            try:
                page = self.get_thread(thread_uid, page_size, page_id)
                self._read_options()
            except SearchIncomplete:
                raise incomplete("Tidsbudsjettet er brukt opp.") from None
            except AuthError:
                raise incomplete("Autentisering feila.") from None
            except httpx.HTTPStatusError as error:
                raise incomplete(f"Vigilo svarte HTTP {error.response.status_code}.") from None
            except httpx.TransportError:
                raise incomplete("Lesekallet feila eller nådde tidsgrensa.") from None
            except ValueError:
                raise incomplete("Ugyldig JSON-respons frå Vigilo.") from None
            pages_fetched += 1
            if (not isinstance(page, dict) or page.get("threadUid") != thread_uid
                    or not isinstance(page.get("messages"), list)):
                raise incomplete("Ugyldig trådside.")
            messages = page["messages"]
            if not messages:
                return {"threadUid": thread_uid, "query": query, "matches": matches,
                        "pagesFetched": pages_fetched, "messagesScanned": messages_scanned,
                        "complete": True}
            returned_page = page.get("pageId")
            if (type(returned_page) is not int or returned_page < 0
                    or returned_page in seen_pages
                    or (page_id is not None and returned_page != page_id)):
                raise incomplete("Ugyldig eller gjenteken pageId.")
            seen_pages.add(returned_page)
            new_messages = 0
            for message in messages:
                if (not isinstance(message, dict) or not isinstance(message.get("id"), str)
                        or not message["id"] or not isinstance(message.get("body"), str)):
                    raise incomplete("Ugyldig melding.")
                timestamp = None
                if since_cutoff is not None:
                    try:
                        timestamp = _message_time(message.get("timeCreated"))
                    except (ValueError, OverflowError):
                        raise incomplete("Ugyldig eller manglande timeCreated i melding.") from None
                if message["id"] in seen_messages:
                    continue
                seen_messages.add(message["id"])
                new_messages += 1
                if since_cutoff is not None and timestamp < since_cutoff:
                    continue
                messages_scanned += 1
                attachments = message.get("attachments") or []
                if (not isinstance(attachments, list)
                        or any(not isinstance(attachment, dict) for attachment in attachments)):
                    raise incomplete("Ugyldige vedlegg.")
                fields = [message["body"], *(attachment.get("name") or "" for attachment in attachments)]
                if any(isinstance(field, str) and query_folded in field.casefold() for field in fields):
                    matches.append({"pageId": returned_page, "message": message})
            if not new_messages:
                raise incomplete("Meldingssida blir gjenteken utan framdrift.")
            page_id = returned_page + 1
        raise incomplete("Sidegrensa er nådd.")


    def message_participants(self, thread_uid: str, child_id: str) -> dict:
        """Deltakere og trådstarter fra appens ordinære meldingstråd."""
        for field, value in (("thread_uid", thread_uid), ("child_id", child_id)):
            if not isinstance(value, str) or not value.strip() or value in (".", ".."):
                raise ValueError(f"{field} må være en gyldig ID.")
        return self.get_json(
            f"/api/message-threads/{quote(thread_uid, safe='')}/participants",
            params={"childPersonId": child_id},
        )

    # --- Web-foreldreportal ---------------------------------------------
    # Endepunkter kartlagt fra web-portalen (web-parent.prod.vigilo-oas.no),
    # 2026-09-12. Dekker det app-gatewayen mangler: timeplan, fravær,
    # samtykke, vurdering, nyheter. Alle bruker childId (=personId) og de
    # fleste også organizationalUnitId (skole/klasse) — begge kommer fra
    # web_children(). `week` er ISO-format "YYYY-WW" (se iso_week()).

    def web_children(self) -> list:
        """Barn fra web-portalen (rå). Hvert barn har `id` + `organizationalUnits`."""
        data = self.web_get_json("/api/children/my")
        return data if isinstance(data, list) else data.get("items", data.get("children", []))

    def children(self) -> list:
        """Én rad per barn/enhet, inkludert SFO, med stabile ID-er og type.

        Rå enhetsfelt og alle grupper beholdes. Barn uten enhet beholdes også.
        """
        out = []
        for c in self.web_children():
            for unit in c.get("organizationalUnits") or [{}]:
                groups = unit.get("groups") or []
                out.append({
                    "childId": c.get("id") or c.get("childId"),
                    "firstName": c.get("firstName"), "lastName": c.get("lastName"),
                    "organizationalUnitId": unit.get("id"),
                    "organizationalUnitType": unit.get("type"),
                    "organizationalUnit": dict(unit),
                    "school": unit.get("name"),
                    "group": groups[0].get("name") if groups else None,
                    "groups": groups,
                })
        return out

    def web_current_user(self) -> dict:
        return self.web_get_json("/api/current-user")

    def news_feed(self, child_id: str, from_date: str | None = None, to_date: str | None = None) -> list:
        """Oppslag/«Siste nytt» for et barn. Datoer på formatet YYYY-MM-DD."""
        from_date = from_date or (date.today() - timedelta(days=30)).isoformat()
        to_date = to_date or (date.today() + timedelta(days=1)).isoformat()
        data = self.web_get_json(
            "/api/news-feed",
            params={"childIds": child_id, "fromDate": from_date, "toDate": to_date},
        )
        return data if isinstance(data, list) else data.get("items", [])

    def web_message_threads(
        self, child_id: str, from_date: str | None = None, to_date: str | None = None
    ) -> list:
        """Beskjedtråder via web-portalen (parallell til app-ens get_message_threads)."""
        from_date = from_date or (date.today() - timedelta(days=30)).isoformat()
        to_date = to_date or (date.today() + timedelta(days=1)).isoformat()
        return self.web_get_json(
            "/api/message-threads",
            params={
                "childIds": child_id,
                "fromDate": from_date,
                "toDate": to_date,
                "includeAfterSchoolMessages": "false",
            },
        )

    def absences(self, child_id: str, include_after_school: bool = False) -> list:
        """Fravær for et barn."""
        return self.web_get_json(
            "/api/absences",
            params={
                "childIds": child_id,
                "includeAfterSchoolAbsences": str(include_after_school).lower(),
            },
        )

    def consent_forms(self, child_id: str, organizational_unit_id: str) -> list:
        """Samtykkeskjemaer for et barn."""
        return self.web_get_json(
            "/api/consent-forms",
            params={"childId": child_id, "organizationalUnitId": organizational_unit_id},
        )

    def lessons(self, child_id: str, organizational_unit_id: str, week: str | None = None) -> list:
        """Timeplan/ukesplan for en gitt ISO-uke (default inneværende uke)."""
        return self.web_get_json(
            f"/api/students/{child_id}/lessons",
            params={"organizationalUnitId": organizational_unit_id, "week": week or iso_week()},
        )

    def scheduling_events(
        self, child_id: str, organizational_unit_id: str, week: str | None = None
    ) -> list:
        """Timeplan-hendelser (prøver, aktiviteter o.l.) for en ISO-uke."""
        return self.web_get_json(
            f"/api/scheduling-events/{child_id}/student",
            params={"organizationalUnitId": organizational_unit_id, "week": week or iso_week()},
        )

    def school_years(self, organizational_unit_id: str) -> list:
        """Skoleår for en enhet — trengs bl.a. for å hente vurderinger per termin."""
        return self.web_get_json(
            "/api/school-years", params={"organizationalUnitId": organizational_unit_id}
        )

    # TODO(kartlegging): selve karakter-/vurderings-endepunktet laster først
    # etter valg av skoleår/termin i Vurdering-fanen — sniff det med web_api_get.

    # --- Vedlegg (f.eks. ukeplan-PDF i en melding) ----------------------
    # Meldinger fra app-gatewayen (get_thread) har vedlegg med en ferdig-signert
    # blob-URL (Azure SAS, gyldig ~6 t) — kan lastes ned direkte uten auth.
    # Hent tråden fersk rett før nedlasting, ellers kan URL-en være utløpt.

    def read_thread_attachments(self, thread_uid: str, max_chars: int = 40000) -> list:
        """Last ned og les vedleggene i en meldingstråd.

        Returnerer én post per vedlegg: {name, mimeType, text} for PDF/tekst, og
        {name, mimeType, note, url} for binærformater vi ikke leser. Brukes bl.a.
        til å lese ukeplan-PDF-er lagt ved i foreldremeldinger.
        """
        thread = self.get_thread(thread_uid)
        out = []
        for msg in thread.get("messages", []):
            for att in msg.get("attachments", []):
                name = att.get("name")
                mime = att.get("mimeType", "")
                url = att.get("url")
                if not url:
                    out.append({"name": name, "mimeType": mime, "note": "mangler URL"})
                    continue
                try:
                    resp = logged_get(url, origin="attachment", timeout=60, follow_redirects=True)
                    resp.raise_for_status()
                    text = _attachment_to_text(resp.content, mime, name)
                    if text is None:
                        out.append(
                            {
                                "name": name,
                                "mimeType": mime,
                                "note": "ikke-tekstlig format — ikke lest",
                                "url": url,
                            }
                        )
                    else:
                        out.append({"name": name, "mimeType": mime, "text": text[:max_chars]})
                except Exception as e:  # noqa: BLE001 — surface, ikke krasj
                    out.append(
                        {"name": name, "mimeType": mime, "note": f"nedlasting/lesing feilet: {e}"}
                    )
        return out

    def read_post_attachments(self, post_id: str, max_chars: int = 40000) -> list:
        """Last ned og les vedleggene i et oppslag/nyhetspost (news_feed).

        Ukeplanen legges typisk ut her som en fler-ukers PDF (f.eks. «uke 36-38»),
        ikke som en meldingstråd. Vedleggene lastes ned fra web-API-et med auth
        (ikke SAS-URL som i meldinger). `post_id` er `id`-feltet fra `news_feed`.
        """
        post = self.web_get_json(f"/api/posts/{post_id}")
        out = []
        for att in post.get("attachments", []):
            name = att.get("name")
            mime = att.get("mimeType", "")
            aid = att.get("id")
            if not aid:
                out.append({"name": name, "mimeType": mime, "note": "mangler id"})
                continue
            try:
                r = self.web_get(f"/api/posts/{post_id}/attachments/{aid}")
                r.raise_for_status()
                text = _attachment_to_text(r.content, mime, name)
                if text is None:
                    out.append(
                        {"name": name, "mimeType": mime, "note": "ikke-tekstlig format — ikke lest"}
                    )
                else:
                    out.append({"name": name, "mimeType": mime, "text": text[:max_chars]})
            except Exception as e:  # noqa: BLE001 — surface, ikke krasj
                out.append(
                    {"name": name, "mimeType": mime, "note": f"nedlasting/lesing feilet: {e}"}
                )
        return out

    def close(self) -> None:
        self._http.close()
        self._web.close()
